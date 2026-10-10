from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .models import ReportInput


REPORT_FONT = "Tajawal"


def _clean(value: str) -> str:
    return str(value or "").strip()


def _find_value_cell(table, label: str):
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if _clean(cell.text) == label and index + 1 < len(row.cells):
                return row.cells[index + 1]
    return None


def _set_cell_text(cell, value: str, alignment=1):
    if cell is None:
        return
    # Preserve the formatting of the first available run.
    first_run = None
    for paragraph in cell.paragraphs:
        if paragraph.runs:
            first_run = paragraph.runs[0]
            break
    font = None
    if first_run is not None:
        font = {
            "name": first_run.font.name,
            "size": first_run.font.size,
            "bold": first_run.bold,
            "italic": first_run.italic,
            "color": first_run.font.color.rgb,
        }
    previous_alignment = cell.paragraphs[0].alignment if cell.paragraphs else None
    cell.text = str(value)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = previous_alignment if alignment is None else alignment
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    if paragraph.runs and font:
        run = paragraph.runs[0]
        run.font.name = font["name"]
        run.font.size = font["size"]
        run.bold = font["bold"]
        run.italic = font["italic"]
        if font["color"] is not None:
            run.font.color.rgb = font["color"]
        r_pr = run._r.get_or_add_rPr()
        rtl = r_pr.find(qn("w:rtl"))
        if rtl is None:
            rtl = OxmlElement("w:rtl")
            r_pr.append(rtl)
        rtl.set(qn("w:val"), "1")


def _apply_report_font(doc: Document):
    """Apply Tajawal to every report style and every existing/generated run."""
    for style in doc.styles:
        if not hasattr(style, "font"):
            continue
        style.font.name = REPORT_FONT
        r_pr = style.font._element.get_or_add_rPr()
        r_fonts = r_pr.rFonts
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            r_fonts.set(qn(f"w:{script}"), REPORT_FONT)

    def apply_paragraph(paragraph):
        for run in paragraph.runs:
            run.font.name = REPORT_FONT
            r_pr = run._r.get_or_add_rPr()
            r_fonts = r_pr.rFonts
            if r_fonts is None:
                r_fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, r_fonts)
            for script in ("ascii", "hAnsi", "eastAsia", "cs"):
                r_fonts.set(qn(f"w:{script}"), REPORT_FONT)

    for paragraph in doc.paragraphs:
        apply_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    apply_paragraph(paragraph)
    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                apply_paragraph(paragraph)


def _find_row(table, label: str):
    for row in table.rows:
        if any(_clean(cell.text) == label for cell in row.cells):
            return row
    return None


def _find_special_notes_table(doc: Document):
    for table in doc.tables:
        if any("ملخص الحالات الخاصة" in _clean(cell.text) for row in table.rows for cell in row.cells):
            return table
    return None


def _delete_row(table, row):
    table._tbl.remove(row._tr)


def _clone_note_row(table, template_row):
    new_tr = deepcopy(template_row._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def _write_special_notes(table, notes: list[str]):
    fixed_index = None
    for index, row in enumerate(table.rows):
        if "رقم الأرض من الحصر الميداني" in _clean(row.cells[0].text):
            fixed_index = index
            break
    if fixed_index is None:
        raise ValueError("لم يتم العثور على ملاحظة رقم الأرض الثابتة داخل جدول الحالات الخاصة.")
    if len(table.rows) <= fixed_index + 1:
        raise ValueError("قالب Word غير متوافق: لا يوجد صف نموذج للملاحظات الإضافية.")

    template_row = table.rows[fixed_index + 1]
    for row in list(table.rows)[fixed_index + 1:]:
        _delete_row(table, row)

    notes = [note.strip(" •") for note in notes if _clean(note) and _clean(note) != "لا يوجد"]
    if not notes:
        notes = ["لا توجد حالات خاصة إضافية وفقًا للبيانات المتاحة."]
    for note in notes:
        row = _clone_note_row(table, template_row)
        _set_cell_text(row.cells[0], note, alignment=None)
        if len(row.cells) > 1:
            _set_cell_text(row.cells[1], "-")


def validate_template(doc: Document):
    required_labels = {
        "الجهة المستلمة", "الجهة المنفذة",
        "اسم الجمعية الرسمي", "نوع الجمعية", "كود الجمعية", "عدد الأحواض",
        "إجمالي عدد القطع", "القطع المسجلة على المنظومة",
        "القطع غير المسجلة على المنظومة", "منها قطع غير محيزة",
    }
    found = set()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = _clean(cell.text)
                if text in required_labels:
                    found.add(text)
    missing = sorted(required_labels - found)
    if missing:
        raise ValueError("قالب Word غير متوافق. الحقول الناقصة: " + "، ".join(missing))
    if not any("ملخص الحالات الخاصة" in _clean(c.text) for t in doc.tables for r in t.rows for c in r.cells):
        raise ValueError("قالب Word غير متوافق: قسم ملخص الحالات الخاصة غير موجود.")


def generate_report(template_path: str | Path, output_path: str | Path, data: ReportInput):
    template = Path(template_path)
    output = Path(output_path)
    if not template.exists():
        raise FileNotFoundError(f"ملف القالب غير موجود: {template}")

    doc = Document(template)
    validate_template(doc)
    _apply_report_font(doc)

    replacements = {
        "الجهة المستلمة": data.recipient,
        "تاريخ التسليم": data.delivery_date,
        "الجهة المنفذة": data.executor,
        "حالة التقرير": data.report_status,
        "اسم الجمعية الرسمي": data.association_name,
        "نوع الجمعية": data.association_type,
        "كود الجمعية": data.association_code,
        "عدد الأحواض": str(data.analysis.basin_count),
    }
    for table in doc.tables:
        for label, value in replacements.items():
            cell = _find_value_cell(table, label)
            if cell is not None:
                _set_cell_text(cell, value)

    stats = {
        "إجمالي عدد القطع": (data.analysis.total_parcels, 100.0),
        "القطع المسجلة على المنظومة": (data.analysis.registered_count, data.analysis.registered_percentage),
        "القطع غير المسجلة على المنظومة": (data.analysis.unregistered_count, data.analysis.unregistered_percentage),
        "منها قطع غير محيزة": (data.analysis.non_held_count, data.analysis.non_held_percentage),
    }
    for table in doc.tables:
        for label, (count, percentage) in stats.items():
            row = _find_row(table, label)
            if row is None or len(row.cells) < 3:
                continue
            # Template visual order: label, count, percentage in XML order.
            _set_cell_text(row.cells[1], str(count))
            percentage_text = "100%" if percentage == 100 else f"{percentage:.1f}%"
            _set_cell_text(row.cells[2], percentage_text)

    # Preserve the fixed land-number note and add descriptive dynamic notes in separate rows.
    special_notes = []
    if data.analysis.placeholder_national_id_count:
        count = data.analysis.placeholder_national_id_count
        special_notes.append(
            f"توجد {count} قطع مسجلة برقم قومي 11111111111111، وهو رقم غير صالح ولا يمثل رقمًا قوميًّا فعليًّا، كما لا يوجد مصدر أو مرجع معتمد لدى الجمعية للبيانات القومية الخاصة بهذه القطع."
        )
    if data.analysis.non_held_count:
        count = data.analysis.non_held_count
        special_notes.append(
            f"توجد {count} قطع غير محيزة لدى الجمعية، أي لم يتم تسجيلها على حيازة معتمدة ضمن بيانات الجمعية."
        )
    if data.analysis.heirs_without_authorized_count:
        count = data.analysis.heirs_without_authorized_count
        special_notes.append(
            f"توجد {count} حالات حيازة مسجلة بأسماء ورثة دون وجود مفوض عنهم، مما قد يشير إلى عدم حسم ممثل الورثة الحاليين أو احتمال وجود نزاع بينهم، أو عدم استكمال تسجيل الأرض أو نقلها إلى الورثة أو ممثليهم الحاليين."
        )
    excluded_note_prefixes = (
        "الأرض بها جزء بور:", "الأرض بها مبنى:", "الأرض بها مصنع:", "الأرض بها مزرعة:",
        "غير محيز:", "حالات ورثة دون مفوضين مسجلين:", "أرقام قومية غير صالحة أو غير مكتملة:",
    )
    special_notes.extend(
        note for note in data.analysis.detected_notes
        if not note.startswith(excluded_note_prefixes)
    )
    manual_notes = _clean(data.special_notes)
    if manual_notes and manual_notes != "لا يوجد":
        special_notes.extend(line.strip(" •-") for line in manual_notes.splitlines() if line.strip())
    notes_table = _find_special_notes_table(doc)
    if notes_table is not None:
        _write_special_notes(notes_table, special_notes)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = f"تقرير تسليم {data.association_name}"
    doc.save(output)
    return output
