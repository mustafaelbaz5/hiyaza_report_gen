from collections import Counter
from pathlib import Path
import re

from openpyxl import load_workbook

from .models import AnalysisResult
from .text_utils import clean_text, normalized_key, tidy_association_name, infer_association_type, infer_village_name


REQUIRED_HEADERS = {
    "كود القطعة",
    "رقم الحيازة",
    "اسم الحائز",
    "الرقم القومي للحائز",
    "اسم الجمعية",
    "كود الجمعية",
    "اسم الحوض",
}


def _is_invalid_holding(value) -> bool:
    return clean_text(value).lower() in {"", "0", "-1", "none", "null"}


def _is_valid_national_id(value) -> bool:
    text = clean_text(value)
    if text.endswith(".0"):
        text = text[:-2]
    return bool(re.fullmatch(r"\d{14}", text)) and text != "11111111111111" and set(text) != {"0"}


def _is_valid_parcel_id(value) -> bool:
    return clean_text(value) not in {"", "0", "-1"}


def _percent(count: int, total: int) -> float:
    return round((count / total * 100), 1) if total else 0.0


def _contains_heirs_without_authorized(note: str, holder: str) -> bool:
    holder_text = clean_text(holder)
    note_text = normalized_key(note)
    is_heirs = bool(re.search(r"(?:^|\s)\d+\s*ورث(?:ة|ه)\b", holder_text)) or "ورثة" in holder_text
    if not is_heirs:
        return False
    delegate_terms = ("مفوض عنه", "المفوض عنه", "مفوضة عنه", "وكيل عنه", "وكيلة عنه")
    return not any(term in note_text for term in delegate_terms)


def _summarize_notes(note_counter: Counter) -> list[str]:
    groups = {
        "الحيازة غير مسجلة على المنظومة": 0,
        "الأرض بها مبنى": 0,
        "الأرض بها مصنع": 0,
        "الأرض بها مزرعة": 0,
        "الأرض بها جزء بور": 0,
        "مفوض عنه": 0,
        "غير محيز": 0,
    }
    for note, count in note_counter.items():
        compact = note.replace("أ", "ا").replace("إ", "ا").replace("ة", "ه")
        if "غير مسجله علي المنظومه" in compact or "غير مسجله على المنظومه" in compact:
            groups["الحيازة غير مسجلة على المنظومة"] += count
        if "مبني" in compact or "مبنى" in note:
            groups["الأرض بها مبنى"] += count
        if "مصنع" in compact:
            groups["الأرض بها مصنع"] += count
        if "مزرعه" in compact:
            groups["الأرض بها مزرعة"] += count
        if "جزء بور" in compact or "حزء بور" in compact:
            groups["الأرض بها جزء بور"] += count
        if "مفوض عنه" in compact:
            groups["مفوض عنه"] += count
        if "غير محيز" in compact:
            groups["غير محيز"] += count
    return [f"{label}: {count} قطعة" for label, count in groups.items() if count]


def analyze_excel(path: str | Path) -> AnalysisResult:
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"ملف Excel غير موجود: {source}")

    wb = load_workbook(source, read_only=True, data_only=True)
    if not wb.sheetnames:
        raise ValueError("ملف Excel لا يحتوي على أي Sheet.")

    # Use the first non-empty sheet. This matches the one-association delivery workbook.
    ws = None
    for sheet_name in wb.sheetnames:
        candidate = wb[sheet_name]
        if candidate.max_row >= 2 and candidate.max_column >= 1:
            ws = candidate
            break
    if ws is None:
        raise ValueError("لم يتم العثور على بيانات داخل ملف Excel.")

    headers = [normalized_key(c.value) for c in ws[1]]
    index = {header: i for i, header in enumerate(headers) if header}
    missing_headers = sorted(REQUIRED_HEADERS - set(index))
    if missing_headers:
        raise ValueError("الأعمدة المطلوبة غير موجودة: " + "، ".join(missing_headers))

    def value(row, header):
        pos = index.get(header)
        return row[pos] if pos is not None and pos < len(row) else None

    rows = []
    for excel_row in ws.iter_rows(min_row=2, values_only=True):
        if not any(clean_text(v) for v in excel_row):
            continue
        rows.append(excel_row)

    total = len(rows)
    if total == 0:
        raise ValueError("ملف Excel لا يحتوي على سجلات قطع.")

    association_values = [tidy_association_name(value(r, "اسم الجمعية")) for r in rows]
    association_values = [x for x in association_values if x]
    association_name = Counter(association_values).most_common(1)[0][0] if association_values else ""
    association_code_values = [clean_text(value(r, "كود الجمعية")) for r in rows]
    association_code_values = [x for x in association_code_values if x]
    association_code = Counter(association_code_values).most_common(1)[0][0] if association_code_values else ""
    association_type = infer_association_type(association_name, ws.title)
    village_name = infer_village_name(association_name)

    basins = {normalized_key(value(r, "اسم الحوض")) for r in rows if normalized_key(value(r, "اسم الحوض"))}
    basin_codes = {normalized_key(value(r, "كود الحوض")) for r in rows if "كود الحوض" in index and normalized_key(value(r, "كود الحوض"))}

    registered = 0
    non_held = 0
    note_counter = Counter()
    heirs_without_authorized_count = 0
    invalid_national_id_count = 0
    placeholder_national_id_count = 0
    ids = Counter()
    warnings = []

    note_header = "ملاحظات علي الحيازة (من فريق العمل)"
    for row in rows:
        parcel_id = clean_text(value(row, "كود القطعة"))
        holding = value(row, "رقم الحيازة")
        holder = clean_text(value(row, "اسم الحائز"))
        national_id = value(row, "الرقم القومي للحائز")
        if parcel_id:
            ids[parcel_id] += 1
        if _is_invalid_holding(holding):
            non_held += 1
        national_id_text = clean_text(national_id).removesuffix(".0")
        if not _is_valid_national_id(national_id):
            invalid_national_id_count += 1
        if national_id_text == "11111111111111":
            placeholder_national_id_count += 1
        if (_is_valid_parcel_id(parcel_id)
                and not _is_invalid_holding(holding)
                and bool(holder)
                and _is_valid_national_id(national_id)):
            registered += 1
        note = normalized_key(value(row, note_header)) if note_header in index else ""
        if note:
            note_counter[note] += 1
        if _contains_heirs_without_authorized(note, holder):
            heirs_without_authorized_count += 1

    unregistered = total - registered
    duplicate_ids = [item for item, count in ids.items() if count > 1]
    if duplicate_ids:
        warnings.append(f"يوجد {len(duplicate_ids)} كود قطعة مكرر في ملف Excel.")
    if len(association_values) and len(set(association_values)) > 1:
        warnings.append("ملف Excel يحتوي على أكثر من اسم جمعية.")
    if len(association_code_values) and len(set(association_code_values)) > 1:
        warnings.append("ملف Excel يحتوي على أكثر من كود جمعية.")
    if basin_codes and len(basin_codes) != len(basins):
        warnings.append(
            f"عدد أسماء الأحواض ({len(basins)}) يختلف عن عدد أكواد الأحواض ({len(basin_codes)}). "
            "سيُستخدم عدد أسماء الأحواض في التقرير."
        )

    detected_notes = _summarize_notes(note_counter)
    if heirs_without_authorized_count:
        detected_notes.append(f"حالات ورثة دون مفوضين مسجلين: {heirs_without_authorized_count} قطعة")
    if invalid_national_id_count:
        detected_notes.append(f"أرقام قومية غير صالحة أو غير مكتملة: {invalid_national_id_count} قطعة")

    wb.close()

    return AnalysisResult(
        source_file=source,
        sheet_name=ws.title,
        association_name=association_name,
        association_code=association_code,
        association_type=association_type,
        village_name=village_name,
        total_parcels=total,
        basin_count=len(basins),
        registered_count=registered,
        unregistered_count=unregistered,
        non_held_count=non_held,
        registered_percentage=_percent(registered, total),
        unregistered_percentage=_percent(unregistered, total),
        non_held_percentage=_percent(non_held, total),
        heirs_without_authorized_count=heirs_without_authorized_count,
        invalid_national_id_count=invalid_national_id_count,
        placeholder_national_id_count=placeholder_national_id_count,
        detected_notes=detected_notes,
        warnings=warnings,
    )

